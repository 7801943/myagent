import asyncio

from openpyxl import Workbook, load_workbook

from myagent.tools.builtin.document_tools import _document_edit_impl, document_edit, document_read
from myagent.tools.builtin.file_edit import file_edit_table
from myagent.tools.builtin.file_write import file_write
from myagent.tools.builtin.pdf_tools import pdf_read
from myagent.tools.builtin.spreadsheet_tools import spreadsheet_edit, spreadsheet_read
from myagent.tools.manager import ToolManager


def run_tool(coro):
    return asyncio.run(coro)


def test_document_tools_hide_version_but_retain_internal_stale_check(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("alpha\nbeta\n", encoding="utf-8")

    read = run_tool(document_read(str(path), start_line=2, end_line=2))
    assert not read.is_error
    assert "beta" in read.content
    version = read.metadata["version"]
    assert version not in read.content

    edited = run_tool(document_edit(
        str(path),
        target_content="beta",
        replacement_content="gamma",
    ))
    assert not edited.is_error, edited.content
    assert path.read_text(encoding="utf-8") == "alpha\ngamma\n"
    assert edited.metadata["previous_version"] == version
    assert edited.metadata["version"] != version

    stale = run_tool(_document_edit_impl(
        path=str(path),
        target_content="gamma",
        replacement_content="delta",
        expected_version=version,
    ))
    assert stale.is_error
    assert "版本不匹配" in stale.content
    assert path.read_text(encoding="utf-8") == "alpha\ngamma\n"


def test_document_edit_preserves_detected_text_encoding(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("中文旧值\n".encode("gb18030"))

    result = run_tool(document_edit(
        str(path),
        target_content="旧值",
        replacement_content="新值",
    ))
    assert not result.is_error, result.content
    assert path.read_bytes().decode("gb18030") == "中文新值\n"


def test_spreadsheet_tools_preview_apply_and_hide_tokens(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Code", "Amount"])
    wb.save(path)

    read = run_tool(spreadsheet_read(
        str(path), sheet_name="Data", cell_range="A1:B2", value_mode="both"
    ))
    assert not read.is_error, read.content
    assert "结构 token:" not in read.content
    assert "内容 token:" not in read.content
    assert read.metadata["structure_token"]
    assert read.metadata["content_token"]
    version = read.metadata["version"]

    preview = run_tool(spreadsheet_edit(
        str(path),
        operation="set_range",
        sheet_name="Data",
        payload={"range": "A2:B2", "values": [["A-1", 10]]},
    ))
    assert not preview.is_error, preview.content
    assert preview.metadata["dry_run"] is True
    assert preview.metadata["version"] == version

    applied = run_tool(spreadsheet_edit(
        str(path),
        operation="set_range",
        sheet_name="Data",
        payload={"range": "A2:B2", "values": [["A-1", 10]]},
        dry_run=False,
    ))
    assert not applied.is_error, applied.content
    assert applied.metadata["version"] != version
    wb = load_workbook(path, data_only=False)
    assert wb["Data"]["A2"].value == "A-1"
    assert wb["Data"]["B2"].value == 10
    wb.close()


def test_spreadsheet_read_token_matches_single_range_edit_check(tmp_path):
    path = tmp_path / "tokens.xlsx"
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = "old"
    wb.save(path)

    read = run_tool(spreadsheet_read(str(path), sheet_name="Data", cell_range="A1:A1"))
    preview = run_tool(file_edit_table(
        str(path),
        operation="set_range",
        sheet_name="Data",
        payload={"range": "A1:A1", "values": [["new"]]},
        expected_content_token=read.metadata["content_token"],
    ))
    assert not preview.is_error, preview.content


def test_pdf_read_fallback_keeps_renderer_page_metadata(tmp_path, monkeypatch):
    fitz = __import__("fitz")
    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "hello")
    doc.save(path)
    doc.close()

    from myagent.tools.builtin import _file_common
    from myagent.tools.builtin.pdf_router import PdfRoutingResult

    monkeypatch.setattr(
        _file_common,
        "parse_pdf_with_routing",
        lambda *_args, **_kwargs: PdfRoutingResult(
            route="image",
            pdf_type="unknown",
            confidence=0.0,
            threshold=0.85,
            page_count=0,
            selected_pages=[],
            reason="parser unavailable",
        ),
    )
    result = run_tool(pdf_read(str(path)))
    assert not result.is_error, result.content
    assert result.metadata["page_count"] == 1
    assert result.metadata["selected_pages"] == [1]
    assert result.metadata["fallback_pages"] == [1]
    assert result.metadata["version"] not in result.content


def test_document_edit_uses_line_only_for_ambiguous_matches(tmp_path):
    path = tmp_path / "lines.txt"
    path.write_text("same\nother\nsame\n", encoding="utf-8")

    mismatch = run_tool(document_edit(
        str(path),
        target_content="same",
        replacement_content="changed",
        line_no=2,
    ))
    assert mismatch.is_error
    assert "不匹配" in mismatch.content
    assert path.read_text(encoding="utf-8") == "same\nother\nsame\n"

    edited = run_tool(document_edit(
        str(path),
        target_content="same",
        replacement_content="changed",
        line_no=3,
    ))
    assert not edited.is_error, edited.content
    assert path.read_text(encoding="utf-8") == "same\nother\nchanged\n"

    # 全文已经唯一匹配时，即使 line_no 不对应也忽略 line_no。
    unique = run_tool(document_edit(
        str(path),
        target_content="other",
        replacement_content="unique",
        line_no=999,
    ))
    assert not unique.is_error, unique.content
    assert path.read_text(encoding="utf-8") == "same\nunique\nchanged\n"


def test_document_edit_exposes_highlight_and_comment_for_docx(tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    path = tmp_path / "annotated.docx"
    doc = Document()
    doc.add_paragraph("old text")
    doc.save(path)

    result = run_tool(document_edit(
        str(path),
        target_content="old text",
        replacement_content="new text",
        highlight="yellow",
        comment="reviewed",
    ))
    assert not result.is_error, result.content

    updated = Document(path)
    assert any(
        run.text == "new text" and run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        for run in updated.paragraphs[0].runs
    )
    assert result.metadata["highlight"] == "yellow"
    assert result.metadata["comment_added"] is True


def test_document_edit_line_no_disambiguates_docx_table_match(tmp_path):
    from docx import Document

    path = tmp_path / "ambiguous.docx"
    doc = Document()
    doc.add_paragraph("same")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "same"
    doc.save(path)

    result = run_tool(document_edit(
        str(path),
        target_content="same",
        replacement_content="table value",
        line_no=3,
    ))
    assert not result.is_error, result.content

    updated = Document(path)
    assert updated.paragraphs[0].text == "same"
    assert updated.tables[0].cell(0, 0).text == "table value"


def test_spreadsheet_edit_applies_highlight_and_comment_to_affected_cells(tmp_path):
    path = tmp_path / "annotated.xlsx"
    wb = Workbook()
    wb.active.title = "Data"
    wb.save(path)

    result = run_tool(spreadsheet_edit(
        str(path),
        operation="set_range",
        sheet_name="Data",
        payload={"range": "A1:B1", "values": [["A", "B"]]},
        dry_run=False,
        highlight="green",
        comment="reviewed",
    ))
    assert not result.is_error, result.content

    wb = load_workbook(path, data_only=False)
    ws = wb["Data"]
    for cell in (ws["A1"], ws["B1"]):
        assert cell.fill.fgColor.rgb.endswith("E2F0D9")
        assert cell.comment.text == "reviewed"
    wb.close()


def test_file_write_rejects_structured_formats(tmp_path):
    path = tmp_path / "fake.docx"
    result = run_tool(file_write(str(path), "not a zip package"))
    assert result.is_error
    assert not path.exists()


def test_format_specific_tools_are_registered_with_compact_schema():
    manager = ToolManager()
    expected = {
        "document_read", "document_edit", "pdf_read", "spreadsheet_read", "spreadsheet_edit",
    }
    assert expected.issubset(set(manager.tool_names))

    schema = manager.get("spreadsheet_edit").parameters_schema
    assert schema["properties"]["operation"]["enum"] == [
        "set_range", "update_cells", "clear_range", "append_rows",
    ]
    assert "oneOf" in schema["properties"]["payload"]
    assert set(schema["properties"]) == {
        "path", "sheet_name", "operation", "payload", "dry_run", "highlight", "comment",
    }
    document_schema = manager.get("document_edit").parameters_schema
    assert set(document_schema["properties"]) == {
        "path", "target_content", "replacement_content", "line_no", "highlight", "comment",
    }
