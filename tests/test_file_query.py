import asyncio
import importlib
import json

import pytest
from docx import Document
from openpyxl import Workbook

file_query_module = importlib.import_module("myagent.tools.builtin.file_query")
from myagent.tools.builtin.file_query import file_query
from myagent.tools.manager import ToolManager


def run_tool(coro):
    return asyncio.run(coro)


def make_fake_response(start_line, end_line, answer='found answer', has_answer=True):
    if not has_answer:
        return json.dumps({'has_answer': False, 'answer': '', 'evidence': []})
    return json.dumps({
        'has_answer': True,
        'answer': answer,
        'evidence': [{'start_line': start_line, 'end_line': end_line, 'reason': 'contains the answer'}],
    })


def patch_subagent(monkeypatch, response):
    async def fake_call(messages, *args, **kwargs):
        return response

    monkeypatch.setattr(file_query_module, '_call_subagent_model', fake_call)


def test_file_query_text_returns_answer_and_original_excerpt(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('intro\nneedle answer\nmore evidence\noutro\n', encoding='utf-8')
    patch_subagent(monkeypatch, make_fake_response(2, 3, answer='The answer is in the needle lines.'))

    result = run_tool(file_query(str(path), 'where is the answer?'))

    assert not result.is_error, result.content
    assert result.metadata['evidence_count'] == 1
    assert result.metadata['evidence'][0]['locator'] == 'L2-L3'
    assert '- The answer is in the needle lines. [S1]' in result.content
    assert '1 | intro' in result.content
    assert '2 | needle answer' in result.content
    assert '3 | more evidence' in result.content


def test_file_query_evidence_mode_keeps_evidence_without_answer_summary(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta evidence\n', encoding='utf-8')
    patch_subagent(monkeypatch, make_fake_response(2, 2, answer='Beta is relevant.'))

    result = run_tool(file_query(str(path), 'beta?', mode='evidence'))

    assert not result.is_error, result.content
    assert 'mode=evidence' in result.content
    assert '[S1] L2-L2' in result.content
    assert result.metadata['evidence'][0]['answer'] == 'Beta is relevant.'


def test_file_query_treats_null_optional_args_as_defaults(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta evidence\n', encoding='utf-8')
    patch_subagent(monkeypatch, make_fake_response(2, 2, answer='Beta default answer.'))

    result = run_tool(file_query(str(path), 'beta?', mode=None, sheet_name=None, max_evidence=None))

    assert not result.is_error, result.content
    assert result.metadata['mode'] == 'answer'
    assert result.metadata['evidence_count'] == 1
    assert 'Beta default answer. [S1]' in result.content


def test_file_query_discards_answer_when_no_evidence(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta\n', encoding='utf-8')
    patch_subagent(monkeypatch, make_fake_response(1, 1, has_answer=False))

    result = run_tool(file_query(str(path), 'missing?'))

    assert not result.is_error, result.content
    assert '未在文件中找到足够依据' in result.content
    assert result.metadata['evidence_count'] == 0


def test_file_query_filters_out_of_chunk_evidence(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta\n', encoding='utf-8')
    patch_subagent(monkeypatch, make_fake_response(99, 100, answer='invalid'))

    result = run_tool(file_query(str(path), 'beta?'))

    assert not result.is_error, result.content
    assert result.metadata['evidence_count'] == 0
    assert 'invalid' not in result.content


def test_file_query_retries_invalid_json_once(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta\n', encoding='utf-8')
    calls = {'count': 0}

    async def fake_call(messages, *args, **kwargs):
        calls['count'] += 1
        if calls['count'] == 1:
            return 'not-json'
        return make_fake_response(2, 2, answer='Recovered answer.')

    monkeypatch.setattr(file_query_module, '_call_subagent_model', fake_call)

    result = run_tool(file_query(str(path), 'beta?'))

    assert not result.is_error, result.content
    assert calls['count'] == 2
    assert 'Recovered answer. [S1]' in result.content


def test_file_query_merges_duplicate_evidence_from_overlapping_chunks(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta\ngamma\n', encoding='utf-8')
    patch_subagent(monkeypatch, make_fake_response(2, 2, answer='Beta answer.'))

    original_chunk_lines = file_query_module._chunk_lines

    def duplicate_chunks(lines, token_limit, overlap_lines):
        return [
            file_query_module.LineChunk(index=1, lines=list(lines)),
            file_query_module.LineChunk(index=2, lines=list(lines)),
        ]

    monkeypatch.setattr(file_query_module, '_chunk_lines', duplicate_chunks)
    result = run_tool(file_query(str(path), 'beta?'))
    monkeypatch.setattr(file_query_module, '_chunk_lines', original_chunk_lines)

    assert not result.is_error, result.content
    assert result.metadata['chunk_count'] == 2
    assert result.metadata['evidence_count'] == 1


def test_file_query_docx_uses_extracted_line_numbers(tmp_path, monkeypatch):
    path = tmp_path / 'doc.docx'
    doc = Document()
    doc.add_paragraph('Doc answer')
    doc.save(path)
    patch_subagent(monkeypatch, make_fake_response(1, 1, answer='DOCX answer.'))

    result = run_tool(file_query(str(path), 'doc?'))

    assert not result.is_error, result.content
    assert result.metadata['format'] == 'docx'
    assert result.metadata['evidence'][0]['locator'] == 'L1-L1'
    assert 'Doc answer' in result.content


def test_file_query_xlsx_locator_includes_sheet_name(tmp_path, monkeypatch):
    path = tmp_path / 'book.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append(['Name', 'Value'])
    ws.append(['Answer', '42'])
    wb.save(path)
    patch_subagent(monkeypatch, make_fake_response(2, 2, answer='XLSX answer.'))

    result = run_tool(file_query(str(path), 'value?', sheet_name='Data'))

    assert not result.is_error, result.content
    assert result.metadata['format'] == 'xlsx'
    assert result.metadata['evidence'][0]['locator'] == 'Data: L2-L2'
    assert '2 | Answer | 42' in result.content


def test_file_query_pdf_locator_includes_page_number(tmp_path, monkeypatch):
    fitz = pytest.importorskip('fitz')
    path = tmp_path / 'sample.pdf'
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), 'PDF answer line')
    doc.save(path)
    doc.close()
    patch_subagent(monkeypatch, make_fake_response(1, 1, answer='PDF answer.'))

    result = run_tool(file_query(str(path), 'pdf?'))

    assert not result.is_error, result.content
    assert result.metadata['format'] == 'pdf'
    assert result.metadata['evidence'][0]['locator'].startswith('page 1')
    assert 'PDF answer line' in result.content


def test_file_query_returns_error_when_subagent_unavailable(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta\n', encoding='utf-8')

    async def failing_call(messages, *args, **kwargs):
        raise RuntimeError('provider unavailable')

    monkeypatch.setattr(file_query_module, '_call_subagent_model', failing_call)

    result = run_tool(file_query(str(path), 'beta?'))

    assert result.is_error
    assert '子代理不可用' in result.content
    assert result.metadata['chunk_failures'][0]['kind'] == 'subagent_error'


def test_file_query_returns_error_when_all_chunks_fail_with_mixed_failure_kinds(tmp_path, monkeypatch):
    path = tmp_path / 'notes.txt'
    path.write_text('alpha\nbeta\n', encoding='utf-8')

    def two_chunks(lines, token_limit, overlap_lines):
        return [
            file_query_module.LineChunk(index=1, lines=[lines[0]]),
            file_query_module.LineChunk(index=2, lines=[lines[1]]),
        ]

    async def fake_query_chunk_with_retry(query, chunk, router=None):
        if chunk.index == 1:
            raise RuntimeError('provider unavailable')
        raise ValueError('响应中未找到 JSON 对象')

    monkeypatch.setattr(file_query_module, '_chunk_lines', two_chunks)
    monkeypatch.setattr(file_query_module, '_query_chunk_with_retry', fake_query_chunk_with_retry)
    monkeypatch.setattr(file_query_module, '_build_subagent_router', lambda: object())

    result = run_tool(file_query(str(path), 'beta?'))

    assert result.is_error
    assert '所有 chunk 均调用失败或返回无效结果' in result.content
    assert result.metadata['failure_summary'] == {'subagent_error': 1, 'invalid_json': 1}


def test_hot_reload_dataclass_tool_registers_without_sys_modules_error(tmp_path):
    tools_dir = tmp_path / 'tools'
    tool_dir = tools_dir / 'dataclass_tool'
    tool_dir.mkdir(parents=True)
    (tool_dir / 'dataclass_tool.py').write_text(
        """from dataclasses import dataclass
from myagent.tools.api import ToolResult

@dataclass
class Payload:
    value: str = 'ok'

async def dataclass_probe() -> ToolResult:
    return ToolResult(content=Payload().value)
""",
        encoding='utf-8',
    )

    manager = ToolManager(tools_dir=str(tools_dir))
    run_tool(manager._scan())

    assert 'dataclass_probe' in manager.tool_names
    record = manager.get('dataclass_probe')
    assert record is not None
    assert record.file_path is not None


def test_file_query_registered_as_builtin_tool():
    manager = ToolManager()
    manager._register_builtin_tools()

    assert 'file_query' in manager.tool_names
    record = manager.get('file_query')
    assert record is not None
    assert record.meta.timeout == 180
    assert record.parameters_schema['required'] == ['path', 'query']
    assert 'query' in record.parameters_schema['properties']
    assert 'overlap_lines' not in record.parameters_schema['properties']
    assert 'excerpt_context_lines' not in record.parameters_schema['properties']
